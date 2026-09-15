"""Tests der Docx-Erzeugung: entsteht eine gültige, korrekt gegliederte Datei?"""

from __future__ import annotations

import zipfile

from docx import Document

from rstrainer import docx_export


def _seitenumbrueche(dokument: Document) -> int:
    return dokument.element.xml.count('w:type="page"')


def _volltext(dokument: Document) -> str:
    teile = [a.text for a in dokument.paragraphs]
    for tabelle in dokument.tables:
        for zeile in tabelle.rows:
            teile.extend(zelle.text for zelle in zeile.cells)
    return "\n".join(teile)


# --- Übungsblatt ------------------------------------------------------------

def test_uebungsblatt_ist_gueltige_docx(tmp_path, register):
    pfad = docx_export.uebungsblatt_schreiben(
        tmp_path / "ub.docx", "TK", "Doppelkonsonanten", ["07"], register,
        "Aufgabe 1: ko___en", "Test 1: Schreibe richtig.",
    )
    assert pfad.exists() and pfad.stat().st_size > 0
    assert zipfile.is_zipfile(pfad)          # .docx ist ein ZIP-Container
    Document(str(pfad))                      # lässt sich wieder öffnen


def test_uebungsblatt_trennt_vorder_und_rueckseite(tmp_path, register):
    """Vorderseite Übung, Rückseite Test – ohne Seitenumbruch stimmt der
    beidseitige Druck nicht."""
    pfad = docx_export.uebungsblatt_schreiben(
        tmp_path / "ub.docx", "TK", "Titel", ["07"], register,
        "ÜBUNGSINHALT", "TESTINHALT",
    )
    dokument = Document(str(pfad))
    assert _seitenumbrueche(dokument) == 1


def test_uebungsblatt_enthaelt_name_datum_und_kategorie(tmp_path, register):
    pfad = docx_export.uebungsblatt_schreiben(
        tmp_path / "ub.docx", "TK", "Kürzemarkierung", ["07"], register,
        "Übung", "Test", datum="05.03.2026",
    )
    text = _volltext(Document(str(pfad)))
    assert "TK" in text
    assert "05.03.2026" in text
    assert "Kürzemarkierung" in text
    assert "07 – Einfachschreibung für Konsonantenverdoppelung" in text


def test_uebungsblatt_enthaelt_beide_inhalte(tmp_path, register):
    pfad = docx_export.uebungsblatt_schreiben(
        tmp_path / "ub.docx", "TK", "Titel", ["07"], register,
        "EINDEUTIGER_UEBUNGSTEXT", "EINDEUTIGER_TESTTEXT",
    )
    text = _volltext(Document(str(pfad)))
    assert "EINDEUTIGER_UEBUNGSTEXT" in text
    assert "EINDEUTIGER_TESTTEXT" in text


def test_loesungen_standardmaessig_nicht_im_dokument(tmp_path, register):
    """Wichtigster Schutz: Lösungen dürfen nicht versehentlich mitgedruckt werden."""
    pfad = docx_export.uebungsblatt_schreiben(
        tmp_path / "ub.docx", "TK", "Titel", ["07"], register,
        "Übung", "Test", loesungen="GEHEIME_LOESUNG",
    )
    assert "GEHEIME_LOESUNG" not in _volltext(Document(str(pfad)))


def test_loesungsblatt_nur_auf_ausdruecklichen_wunsch(tmp_path, register):
    pfad = docx_export.uebungsblatt_schreiben(
        tmp_path / "ub.docx", "TK", "Titel", ["07"], register,
        "Übung", "Test", loesungen="GEHEIME_LOESUNG", loesungen_anhaengen=True,
    )
    dokument = Document(str(pfad))
    text = _volltext(dokument)
    assert "GEHEIME_LOESUNG" in text
    assert "LÖSUNGSBLATT – NICHT AUSTEILEN" in text
    assert _seitenumbrueche(dokument) == 2


def test_leeres_loesungsfeld_erzeugt_keine_dritte_seite(tmp_path, register):
    pfad = docx_export.uebungsblatt_schreiben(
        tmp_path / "ub.docx", "TK", "Titel", ["07"], register,
        "Übung", "Test", loesungen="   ", loesungen_anhaengen=True,
    )
    assert _seitenumbrueche(Document(str(pfad))) == 1


def test_mehrzeiliger_inhalt_bleibt_erhalten(tmp_path, register):
    inhalt = "Zeile eins\nZeile zwei\n\nZeile vier"
    pfad = docx_export.uebungsblatt_schreiben(
        tmp_path / "ub.docx", "TK", "Titel", ["07"], register, inhalt, "Test",
    )
    text = _volltext(Document(str(pfad)))
    for zeile in ["Zeile eins", "Zeile zwei", "Zeile vier"]:
        assert zeile in text


def test_verzeichnis_wird_bei_bedarf_angelegt(tmp_path, register):
    pfad = docx_export.uebungsblatt_schreiben(
        tmp_path / "neu" / "tiefer" / "ub.docx", "TK", "Titel", ["07"], register,
        "Übung", "Test",
    )
    assert pfad.exists()


# --- Informationsblatt ------------------------------------------------------

def test_informationsblatt_listet_fehler_in_tabelle(tmp_path, register):
    fehler = [
        {"wort_original": "Hund", "wort_schueler": "Hunt", "kategorie_nr": "19",
         "kontext": "der [Hund] lief"},
        {"wort_original": "kommen", "wort_schueler": "komen", "kategorie_nr": "07",
         "kontext": "sie [kommen] her"},
    ]
    pfad = docx_export.informationsblatt_schreiben(
        tmp_path / "info.docx", "TK", "Wald-Diktat", "2026-03-05", fehler, register,
        {"wortzahl_original": 90, "fehlerquote_prozent": 2.2}, "Kurzkommentar hier.",
    )
    dokument = Document(str(pfad))
    fehlertabelle = dokument.tables[1]
    assert len(fehlertabelle.rows) == 3          # Kopfzeile + 2 Fehler
    text = _volltext(dokument)
    assert "Hunt" in text and "komen" in text
    assert "19 – p, t, k für b, d, g im Silbenrand oder Silbenende" in text
    assert "Kurzkommentar hier." in text
    assert "2.2" in text


def test_informationsblatt_zaehlt_kategorien(tmp_path, register):
    fehler = [{"wort_original": "a", "wort_schueler": "b", "kategorie_nr": "07",
               "kontext": ""} for _ in range(3)]
    pfad = docx_export.informationsblatt_schreiben(
        tmp_path / "info.docx", "TK", "D", "2026-03-05", fehler, register,
    )
    assert "3×" in _volltext(Document(str(pfad)))


def test_informationsblatt_ohne_fehler(tmp_path, register):
    pfad = docx_export.informationsblatt_schreiben(
        tmp_path / "info.docx", "TK", "D", "2026-03-05", [], register,
    )
    assert "Keine Fehler erfasst." in _volltext(Document(str(pfad)))


def test_diktattext_kommt_auf_eigene_seite(tmp_path, register):
    pfad = docx_export.informationsblatt_schreiben(
        tmp_path / "info.docx", "TK", "D", "2026-03-05", [], register,
        diktattext="Der Originaltext.",
    )
    dokument = Document(str(pfad))
    assert _seitenumbrueche(dokument) == 1
    assert "Der Originaltext." in _volltext(dokument)


# --- Verlaufsbericht --------------------------------------------------------

def test_verlaufsbericht_mit_trendzeilen(tmp_path, register):
    zeilen = [{"kategorie_nr": "07", "summe_gesamt": 12, "text": "↓ abnehmend",
               "rate_vorher": 6.0, "rate_aktuell": 2.0}]
    pfad = docx_export.verlaufsbericht_schreiben(
        tmp_path / "bericht.docx", "TK", "Jan–Mär 2026", zeilen, register,
        kommentar="Deutliche Fortschritte.",
    )
    text = _volltext(Document(str(pfad)))
    assert "Jan–Mär 2026" in text
    assert "abnehmend" in text
    assert "6.0 → 2.0" in text
    assert "Deutliche Fortschritte." in text


def test_verlaufsbericht_ohne_daten(tmp_path, register):
    pfad = docx_export.verlaufsbericht_schreiben(
        tmp_path / "bericht.docx", "TK", "–", [], register,
    )
    assert "Noch keine auswertbaren Daten" in _volltext(Document(str(pfad)))


def test_verlaufsbericht_bettet_diagramm_ein(tmp_path, register):
    from rstrainer import charts
    from conftest import diktatpunkte

    figur, _ = charts.verlauf_linien(diktatpunkte(3), {"07": [1, 2, 3]}, register)
    bild = charts.speichern(figur, tmp_path / "diagramm.png")
    pfad = docx_export.verlaufsbericht_schreiben(
        tmp_path / "bericht.docx", "TK", "Zeitraum", [], register, diagramm_pfad=bild,
    )
    with zipfile.ZipFile(pfad) as archiv:
        bilder = [n for n in archiv.namelist() if n.startswith("word/media/")]
    assert bilder, "Diagramm wurde nicht eingebettet"


def test_fehlendes_diagramm_bricht_nicht_ab(tmp_path, register):
    pfad = docx_export.verlaufsbericht_schreiben(
        tmp_path / "bericht.docx", "TK", "Zeitraum", [], register,
        diagramm_pfad=tmp_path / "gibt_es_nicht.png",
    )
    assert pfad.exists()
