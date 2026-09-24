"""Word-Export (.docx) für Übungsblätter, Mini-Tests und Informationsblätter.

Layout Übungsblatt
------------------
Seite 1 = Übungsteil, Seite 2 = Mini-Test (harter Seitenumbruch, damit der
beidseitige Druck sicher stimmt). Das Lösungsblatt kommt optional auf eine
dritte Seite und ist deutlich als solches markiert, damit es nicht
versehentlich mit ausgeteilt wird.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt, RGBColor

from .kategorien import Register, schwerpunkte

GRAU = RGBColor(0x55, 0x55, 0x55)
ROT = RGBColor(0xB0, 0x22, 0x22)


def _seite_einrichten(dokument: Document) -> None:
    for abschnitt in dokument.sections:
        abschnitt.top_margin = Cm(2.0)
        abschnitt.bottom_margin = Cm(2.0)
        abschnitt.left_margin = Cm(2.2)
        abschnitt.right_margin = Cm(2.0)
    stil = dokument.styles["Normal"]
    stil.font.name = "Calibri"
    stil.font.size = Pt(12)
    stil.paragraph_format.space_after = Pt(6)


def _kopfzeile(dokument: Document, name: str, datum: str, titel: str,
               untertitel: str = "") -> None:
    """Name / Datum links und rechts, darunter der Titel."""
    tabelle = dokument.add_table(rows=1, cols=2)
    tabelle.autofit = True
    links = tabelle.cell(0, 0).paragraphs[0]
    links.add_run("Name: ").bold = True
    links.add_run(name)
    rechts = tabelle.cell(0, 1).paragraphs[0]
    rechts.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rechts.add_run("Datum: ").bold = True
    rechts.add_run(datum)

    ueberschrift = dokument.add_paragraph()
    ueberschrift.paragraph_format.space_before = Pt(10)
    lauf = ueberschrift.add_run(titel)
    lauf.bold = True
    lauf.font.size = Pt(16)

    if untertitel:
        zeile = dokument.add_paragraph()
        lauf = zeile.add_run(untertitel)
        lauf.italic = True
        lauf.font.size = Pt(10)
        lauf.font.color.rgb = GRAU

    trenner = dokument.add_paragraph()
    trenner.paragraph_format.space_after = Pt(10)
    trenner.add_run("─" * 58).font.color.rgb = GRAU


def _kategorien_untertitel(kategorien: list[str], register: Register) -> str:
    if not kategorien:
        return ""
    return "Förderschwerpunkte: " + "  •  ".join(register.label(nr) for nr in kategorien)


def _textblock(dokument: Document, text: str) -> None:
    """Übernimmt einen mehrzeiligen Text absatzweise, Leerzeilen bleiben erhalten."""
    for zeile in (text or "").split("\n"):
        absatz = dokument.add_paragraph()
        absatz.paragraph_format.space_after = Pt(4)
        if zeile.strip():
            absatz.add_run(zeile.rstrip())


def _seitenumbruch(dokument: Document) -> None:
    dokument.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------------------
# Übungsblatt + Mini-Test
# ---------------------------------------------------------------------------

def uebungsblatt_schreiben(pfad: Path | str, schuelername: str, titel: str,
                           kategorien: list[str], register: Register,
                           inhalt_uebung: str, inhalt_test: str,
                           loesungen: str = "", datum: str | None = None,
                           loesungen_anhaengen: bool = False) -> Path:
    """Schreibt Übungsblatt (Seite 1) und Mini-Test (Seite 2) in eine .docx."""
    pfad = Path(pfad)
    pfad.parent.mkdir(parents=True, exist_ok=True)
    datum = datum or date.today().strftime("%d.%m.%Y")

    dokument = Document()
    _seite_einrichten(dokument)

    # --- Vorderseite: Übungsblatt
    _kopfzeile(dokument, schuelername, datum, titel,
               _kategorien_untertitel(kategorien, register))
    _textblock(dokument, inhalt_uebung)

    # --- Rückseite: Mini-Test
    _seitenumbruch(dokument)
    _kopfzeile(dokument, schuelername, datum, f"Mini-Test: {titel}",
               _kategorien_untertitel(kategorien, register))
    _textblock(dokument, inhalt_test)

    # --- Optional: Lösungsblatt, deutlich markiert
    if loesungen_anhaengen and loesungen.strip():
        _seitenumbruch(dokument)
        warnung = dokument.add_paragraph()
        lauf = warnung.add_run("LÖSUNGSBLATT – NICHT AUSTEILEN")
        lauf.bold = True
        lauf.font.size = Pt(14)
        lauf.font.color.rgb = ROT
        _kopfzeile(dokument, schuelername, datum, f"Lösungen: {titel}")
        _textblock(dokument, loesungen)

    dokument.save(str(pfad))
    return pfad


# ---------------------------------------------------------------------------
# Informationsblatt zu einem Diktat
# ---------------------------------------------------------------------------

def informationsblatt_schreiben(pfad: Path | str, schuelername: str,
                                diktat_titel: str, datum: str,
                                fehler: list[dict], register: Register,
                                kennzahlen: dict | None = None,
                                kommentar: str = "",
                                diktattext: str = "",
                                art: str = "diktat") -> Path:
    """Fehlerliste, Schwerpunkte und Kurzkommentar zu einem Text."""
    pfad = Path(pfad)
    pfad.parent.mkdir(parents=True, exist_ok=True)
    bezeichnung = {"freitext": "Freier Text", "diktiert": "Diktierter Text (Sprachsoftware) – nur Satzbau und Grammatik"}.get(art, "Diktat")

    dokument = Document()
    _seite_einrichten(dokument)
    _kopfzeile(dokument, schuelername, datum, f"Auswertung: {diktat_titel}",
               f"Rechtschreibförderung – Informationsblatt · {bezeichnung}")

    if kennzahlen:
        absatz = dokument.add_paragraph()
        absatz.add_run("Überblick: ").bold = True
        absatz.add_run(
            f"{kennzahlen.get('wortzahl_original', 0)} Wörter, "
            f"{len(fehler)} erfasste Fehler"
        )
        if kennzahlen.get("fehlerquote_prozent") is not None:
            absatz.add_run(f", Fehlerquote {kennzahlen['fehlerquote_prozent']} %")
        absatz.add_run(".")

    # --- Schwerpunkte zuerst: Das ist die Antwort auf «Woran arbeiten wir?».
    # Die vollständige Fehlerliste darunter belegt sie nur.
    punkte = schwerpunkte(fehler, hoechstens=8)
    if punkte.liste:
        ueberschrift = dokument.add_paragraph()
        ueberschrift.paragraph_format.space_before = Pt(12)
        ueberschrift.add_run("Schwerpunkte").bold = True
        for nr, anzahl in punkte.liste:
            dokument.add_paragraph(f"{register.label(nr)}: {anzahl}×",
                                   style="List Bullet")
        if punkte.rest:
            dokument.add_paragraph(
                f"Dazu {punkte.rest} weitere Fehlerarten mit einzelnen Vorkommen."
            )

    # --- Fehlerliste als Tabelle
    ueberschrift = dokument.add_paragraph()
    ueberschrift.paragraph_format.space_before = Pt(12)
    ueberschrift.add_run("Erfasste Fehler").bold = True

    if fehler:
        tabelle = dokument.add_table(rows=1, cols=4)
        tabelle.style = "Table Grid"
        kopf = tabelle.rows[0].cells
        for spalte, beschriftung in enumerate(
                ["Richtig", "Geschrieben", "Fehlerart", "Im Text"]):
            kopf[spalte].paragraphs[0].add_run(beschriftung).bold = True
        for eintrag in fehler:
            zeile = tabelle.add_row().cells
            zeile[0].text = str(eintrag.get("wort_original", ""))
            zeile[1].text = str(eintrag.get("wort_schueler", ""))
            zeile[2].text = register.label(str(eintrag.get("kategorie_nr", "")))
            zeile[3].text = str(eintrag.get("kontext", ""))
    else:
        dokument.add_paragraph("Keine Fehler erfasst.")

    if kommentar.strip():
        ueberschrift = dokument.add_paragraph()
        ueberschrift.paragraph_format.space_before = Pt(12)
        ueberschrift.add_run("Kommentar").bold = True
        _textblock(dokument, kommentar)

    if diktattext.strip():
        _seitenumbruch(dokument)
        ueberschrift = dokument.add_paragraph()
        ueberschrift.add_run(
            "Text des Kindes" if art in ("freitext", "diktiert") else "Diktattext (Original)").bold = True
        _textblock(dokument, diktattext)

    dokument.save(str(pfad))
    return pfad


# ---------------------------------------------------------------------------
# Verlaufsbericht (für Elterngespräche)
# ---------------------------------------------------------------------------

def verlaufsbericht_schreiben(pfad: Path | str, schuelername: str,
                              zeitraum: str, trendzeilen: list[dict],
                              register: Register,
                              diagramm_pfad: Path | str | None = None,
                              kommentar: str = "") -> Path:
    """Übersicht über den Lernverlauf, geeignet zum Ausdrucken."""
    pfad = Path(pfad)
    pfad.parent.mkdir(parents=True, exist_ok=True)

    dokument = Document()
    _seite_einrichten(dokument)
    _kopfzeile(dokument, schuelername, date.today().strftime("%d.%m.%Y"),
               "Lernverlauf Rechtschreibung", f"Betrachteter Zeitraum: {zeitraum}")

    if diagramm_pfad and Path(diagramm_pfad).exists():
        dokument.add_picture(str(diagramm_pfad), width=Cm(16))
        dokument.add_paragraph()

    if trendzeilen:
        tabelle = dokument.add_table(rows=1, cols=4)
        tabelle.style = "Table Grid"
        kopf = tabelle.rows[0].cells
        for spalte, beschriftung in enumerate(
                ["Fehlerart", "Fehler gesamt", "Entwicklung", "Fehler/100 Wörter"]):
            kopf[spalte].paragraphs[0].add_run(beschriftung).bold = True
        for zeile_daten in trendzeilen:
            zeile = tabelle.add_row().cells
            zeile[0].text = register.label(str(zeile_daten.get("kategorie_nr", "")))
            zeile[1].text = str(zeile_daten.get("summe_gesamt", 0))
            zeile[2].text = str(zeile_daten.get("text", ""))
            zeile[3].text = (f"{zeile_daten.get('rate_vorher', 0)} → "
                             f"{zeile_daten.get('rate_aktuell', 0)}")
    else:
        dokument.add_paragraph("Noch keine auswertbaren Daten vorhanden.")

    if kommentar.strip():
        ueberschrift = dokument.add_paragraph()
        ueberschrift.paragraph_format.space_before = Pt(12)
        ueberschrift.add_run("Einschätzung der Lehrperson").bold = True
        _textblock(dokument, kommentar)

    dokument.save(str(pfad))
    return pfad
